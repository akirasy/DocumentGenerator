import customtkinter as ctk
import tkinter.messagebox as messagebox
from tkinter import filedialog
import csv
from pathlib import Path
import dxpdf
from docx import Document

# Set theme and color
ctk.set_appearance_mode("System")
ctk.set_default_color_theme("blue")

class DocumentGeneratorApp(ctk.CTk):
    def __init__(self):
        super().__init__()
        self.title("Document Generator")
        self.geometry("550x420")
        self.resizable(False, False)

        # File paths
        self.template_path = ctk.StringVar()
        self.csv_path = ctk.StringVar()

        self.setup_ui()

    def setup_ui(self):
        # Header
        header = ctk.CTkLabel(self, text="Document Generator", 
                              font=ctk.CTkFont(size=20, weight="bold"))
        header.pack(pady=(20, 15))

        # Main Card Frame
        main_frame = ctk.CTkFrame(self, corner_radius=10)
        main_frame.pack(padx=20, pady=5, fill="both", expand=True)

        # --- Template Selection ---
        template_label = ctk.CTkLabel(main_frame, text="1. Select .docx Template:", font=ctk.CTkFont(weight="bold"))
        template_label.grid(row=0, column=0, padx=20, pady=(15, 0), sticky="w")
        
        template_entry = ctk.CTkEntry(main_frame, textvariable=self.template_path, width=350, state="readonly")
        template_entry.grid(row=1, column=0, padx=20, pady=(5, 10), sticky="w")
        
        template_btn = ctk.CTkButton(main_frame, text="Browse", width=100, command=self.browse_template)
        template_btn.grid(row=1, column=1, padx=(0, 20), pady=(5, 10))

        # --- CSV Selection ---
        csv_label = ctk.CTkLabel(main_frame, text="2. Select .csv Data Source:", font=ctk.CTkFont(weight="bold"))
        csv_label.grid(row=2, column=0, padx=20, pady=(5, 0), sticky="w")
        
        csv_entry = ctk.CTkEntry(main_frame, textvariable=self.csv_path, width=350, state="readonly")
        csv_entry.grid(row=3, column=0, padx=20, pady=(5, 20), sticky="w")
        
        csv_btn = ctk.CTkButton(main_frame, text="Browse", width=100, command=self.browse_csv)
        csv_btn.grid(row=3, column=1, padx=(0, 20), pady=(5, 20))

        # Action Buttons Frame
        button_frame = ctk.CTkFrame(self, fg_color="transparent")
        button_frame.pack(padx=20, pady=20, fill="x")

        # Primary Action
        generate_btn = ctk.CTkButton(button_frame, text="Auto-Generate DOCX & PDF", 
                                     height=40, font=ctk.CTkFont(weight="bold"), 
                                     fg_color="#28a745", hover_color="#218838",
                                     command=self.generate_all)
        generate_btn.pack(fill="x", pady=(0, 10))

        # Secondary Action
        batch_btn = ctk.CTkButton(button_frame, text="Batch Convert DOCX to PDF (Manual Mode)", 
                                  height=35, fg_color="#6c757d", hover_color="#5a6268",
                                  command=self.batch_convert_pdf)
        batch_btn.pack(fill="x")

    def browse_template(self):
        filename = filedialog.askopenfilename(filetypes=[("Word Documents", "*.docx")])
        if filename:
            self.template_path.set(filename)

    def browse_csv(self):
        filename = filedialog.askopenfilename(filetypes=[("CSV Files", "*.csv")])
        if filename:
            self.csv_path.set(filename)

    def replace_text_in_docx(self, doc, row_data):
        for paragraph in doc.paragraphs:
            for key, value in row_data.items():
                if key in paragraph.text:
                    paragraph.text = paragraph.text.replace(key, str(value))
        
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for key, value in row_data.items():
                            if key in paragraph.text:
                                paragraph.text = paragraph.text.replace(key, str(value))

    def generate_all(self):
        template = self.template_path.get()
        csv_file = self.csv_path.get()

        if not all([template, csv_file]):
            messagebox.showerror("Error", "Please select both the template and CSV files before generating.")
            return

        try:
            # Set basedir
            base_dir = Path.cwd()
            output_folder = base_dir / "output"
            
            docx_dir = output_folder / "docx"
            pdf_dir = output_folder / "pdf"
            
            # Create directories
            docx_dir.mkdir(parents=True, exist_ok=True)
            pdf_dir.mkdir(parents=True, exist_ok=True)

            with open(csv_file, mode='r', encoding='utf-8') as file:
                reader = csv.DictReader(file)
                headers = reader.fieldnames
                
                if not headers or len(headers) < 2:
                    messagebox.showerror("Error", "CSV file must have at least two columns to format the filename.")
                    return
                    
                col1 = headers[0]
                col2 = headers[1]
                count = 0

                for row in reader:
                    doc = Document(template)
                    self.replace_text_in_docx(doc, row)
                    
                    doc_name = f"{row[col1]} - {row[col2]}.docx"
                    safe_doc_name = "".join([c for c in doc_name if c.isalpha() or c.isdigit() or c in ' ._-()']).rstrip()
                    
                    docx_save_path = docx_dir / safe_doc_name
                    pdf_save_path = pdf_dir / safe_doc_name.replace(".docx", ".pdf")
                    
                    doc.save(str(docx_save_path))
                    dxpdf.convert_file(str(docx_save_path), str(pdf_save_path))
                    
                    count += 1

            messagebox.showinfo("Success", f"Successfully generated {count} DOCX and PDF pairs in:\n{output_folder}")

        except Exception as e:
            messagebox.showerror("Error", f"An error occurred during generation:\n{str(e)}")

    def batch_convert_pdf(self):
        target_folder = filedialog.askdirectory(title="Select Folder containing DOCX files (e.g., output/docx)")
        
        if not target_folder:
            return
            
        try:
            target_dir = Path(target_folder)
            base_dir = Path.cwd()
            pdf_out_dir = base_dir / "output" / "pdf"
            pdf_out_dir.mkdir(parents=True, exist_ok=True)
            
            count = 0
            
            for file_path in target_dir.iterdir():
                if file_path.is_file() and file_path.suffix.lower() == ".docx" and not file_path.name.startswith("~"):
                    pdf_path = pdf_out_dir / file_path.with_suffix('.pdf').name
                    try:
                        dxpdf.convert_file(str(file_path), str(pdf_path))
                        count += 1
                    except Exception as file_error:
                        print(f"Failed to convert {file_path.name}: {str(file_error)}")

            if count == 0:
                messagebox.showwarning("No Files Converted", f"No valid .docx files were found in:\n{target_dir}\n\nMake sure you are selecting the folder that actually contains the Word documents.")
            else:
                messagebox.showinfo("Success", f"Finished converting {count} DOCX files to PDF in:\n{pdf_out_dir}")
                
        except Exception as e:
             messagebox.showerror("Error", f"An error occurred during batch conversion:\n{str(e)}")

if __name__ == "__main__":
    app = DocumentGeneratorApp()
    app.mainloop()